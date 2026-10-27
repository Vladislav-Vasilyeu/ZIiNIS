"""
Загрузка и сохранение текстовых контейнеров (TXT, DOCX).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document


def load_text(path: str | Path) -> str:
    """Загрузить текст из .txt или .docx."""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"Файл не найден: {path}")

    suffix = path.suffix.lower()
    if suffix == ".txt":
        return path.read_text(encoding="utf-8")
    if suffix == ".docx":
        doc = Document(str(path))
        paragraphs = [p.text for p in doc.paragraphs]
        return "\n".join(paragraphs)
    raise ValueError(f"Неподдерживаемый формат файла: {suffix}. Ожидается .txt или .docx")


def save_text(path: str | Path, text: str) -> None:
    """Сохранить текст в .txt или .docx."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    suffix = path.suffix.lower()
    if suffix == ".txt":
        path.write_text(text, encoding="utf-8")
        return
    if suffix == ".docx":
        doc = Document()
        # Разбиваем по абзацам
        for para in text.split("\n"):
            doc.add_paragraph(para)
        doc.save(str(path))
        return
    raise ValueError(f"Неподдерживаемый формат файла: {suffix}. Ожидается .txt или .docx")
